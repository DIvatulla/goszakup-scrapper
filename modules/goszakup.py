from urllib.parse import quote
from urllib.parse import unquote
import bs4
from bs4 import BeautifulSoup
from openpyxl import Workbook
from urllib.parse import urlsplit
import time
import re
import http.client
import random


class goszakup_filters:
	def __init__(self, name: str="", customer: str="", spec: str="",\
				number: str="", month: list=[], year: list=[], status: str="",\
				subject_type: str="", method: list=[], qvazi: str="",\
				count_record: int=2000, page: int=1):
		self.name = name
		self.customer = customer
		self.number = number
		self.month = month
		self.year = year
		self.status = status
		self.subject_type = subject_type
		self.method = method
		self.qvazi = qvazi
		self.count_record = count_record
		self.page = page

	def __iter__(self):
		for attr, value in self.__dict__.items():
			yield attr, value	 

	def urlify(self) -> str:
		res = "/ru/registry/plan?"
		filters_table = dict(self)

		for item in filters_table:
			if (type(filters_table[item]) == int) or (len(filters_table[item]) == 0):
				continue
			
			if type(filters_table[item]) == list:
				for e in filters_table[item]:
					res += quote("filter[{}][]={}&".format(item, e), safe='/:=&')
			else:
				res += quote("filter[{}]={}&".format(item, filters_table[item]), safe='/:=&')	

		res += "count_record={}&page={}".format(self.count_record, self.page)
		return res

class goszakup:
	def __init__(self, filters: goszakup_filters):
		self.filters = filters
		self.host = "goszakup.gov.kz"
		self.headers = {
			"Host": "goszakup.gov.kz",
			"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/120.0.0.0",
			"Content-Type": "*/*"
		}
		self.__conn = None
		self.connect()
		self.__get_cookies()
		self.__count_pages()

	def __get_cookies(self):
		r = self.__request("GET", "/", self.headers)
		r.read()
		for h in reversed(r.getheaders()):
			if h[0] == "Set-Cookie":
				self.headers.setdefault("Set-Cookie", h[1])
				break

	def connect(self):
		if self.__conn != None:
			self.__conn.close()
		
		self.__conn = http.client.HTTPSConnection(self.host, 443)
		self.__get_cookies()

	def __request(self, method: str, url: str, headers: dict) -> http.client.HTTPResponse:
		self.__conn.request(method, url, headers=headers)
		response = self.__conn.getresponse()
		return response
	
	def get(self) -> str:
		r = self.__request("GET", self.filters.urlify(), self.headers)
		return r.read().decode() 

	def __count_pages(self) -> int:
		self.count = 0
		html_doc = self.get()
		soup = BeautifulSoup(html_doc, features="lxml")
		text = ((soup.find_all("small")[0]).find("strong")).get_text(strip=True)
		match = re.search(f"{r"из "}(.*?){r" записей"}", text)
		self.count = int(match.group(1))
		
class excel:
	@classmethod
	def parse_cell(cls, c) -> dict:
		result = {"content": None, "url": None}

		url_tag = c.find('a')
		if url_tag:
			result["url"] = url_tag.get('href')

		result["content"] = c.get_text(strip=True)

		return result

	@classmethod
	def parse_html(cls, html_doc: str) -> list:
		#print(html_doc)
		table = []
		soup = BeautifulSoup(html_doc, features="lxml")
		#print(soup.find_all('tbody')[1])
		
		for tr in ((soup.find_all('tbody')[1]).find_all('tr')):
			cells = tr.find_all('td')
			row = list(map(lambda cell: excel.parse_cell(cell), cells))
			table.append(row)

		return table

	@classmethod
	def make_table(cls, gz: goszakup, filename: str):
		wb = Workbook()
		ws = wb.active

		ws.append(["#п/п", "Заказчик", "Наименование", "Способ_закупки",\
		"Единица_измерения", "Кол-во", "Цена_за_ед", "Плановая_сумма",\
		"Планируемый_срок_закупки", "Статус"])

		gz.filters.page = 0
		try:
			for i in range(gz.filters.count_record, gz.count+gz.filters.count_record, gz.filters.count_record):
				gz.filters.page += 1
				gz.connect()

				for row in cls.parse_html(gz.get()):
					#print(row)
					ws.append([cell_dict["content"] for cell_dict in row])

					for col_idx, cell_dict in enumerate(row, start=1):
						if cell_dict["url"]:
							cell_obj = ws.cell(row=ws.max_row, column=col_idx)
							cell_obj.hyperlink = cell_dict["url"]
							cell_obj.style = "Hyperlink"
				
				print(gz.filters.page)
				time.sleep(random.uniform(9, 29))
		finally:
			wb.save(filename)
